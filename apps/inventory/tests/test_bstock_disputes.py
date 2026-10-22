"""E2b B-Stock disputes: affected lines in B-Stock's layout, the public page, downloads, filing freezes, view log."""

import io
import zipfile
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from rest_framework.test import APIClient

from apps.inventory.models import (
    BStockDispute,
    BStockDisputeEvent,
    ManifestRow,
    ProcessingRow,
    PurchaseOrder,
    Vendor,
)
from apps.inventory.services.bstock_disputes import affected_lines


class DisputeBase(TestCase):
    def setUp(self):
        self.vendor = Vendor.objects.create(name='Test Seller', code='TSTSL')
        self.client = APIClient()
        self.user = get_user_model().objects.create_user(
            email='mgr@example.com', first_name='Mo', last_name='Mgr', password='pw',
        )
        self.user.groups.add(Group.objects.get_or_create(name='Manager')[0])
        self.client.force_authenticate(user=self.user)
        self.order = PurchaseOrder.objects.create(
            vendor=self.vendor, order_number='TRGET-OXX-1', ordered_date='2026-10-01',
            purchase_cost=Decimal('200.00'), retail_value=Decimal('1000.00'), status='processing',
            description='2 Pallets of Skin Care',
        )
        self.prs = []
        for n, (qty, retail, upc) in enumerate(((10, '50.00', '012345678905'), (5, '100.00', '')), start=1):
            mr = ManifestRow.objects.create(
                purchase_order=self.order, row_number=n, quantity=qty, title=f'Lotion {n}',
                unit_retail=Decimal(retail), condition='good',
                identifiers={'upc': upc} if upc else {}, tracking={'pallet_id': f'PAL-{n}', 'lpn': f'LPN{n}'},
            )
            self.prs.append(ProcessingRow.objects.create(
                purchase_order=self.order, manifest_row=mr, row_number=n, quantity=qty, title=f'Lotion {n}',
                unit_retail=Decimal(retail), shelf_price=Decimal('5.00'),
            ))

    def _check_in(self, idx, qty, condition='good'):
        pr = self.prs[idx]
        r = self.client.post(
            f'/api/inventory/orders/{self.order.id}/processing-row-check-in/',
            {
                'processing_row_id': pr.id, 'quantity': qty, 'condition': condition, 'dispatch': 'on_shelf',
                'price': '5.00', 'product_mode': 'new', 'title': pr.title,
            },
            format='json',
        )
        self.assertEqual(r.status_code, 200, r.data)

    def _start(self):
        r = self.client.post(f'/api/inventory/orders/{self.order.id}/bstock-disputes/', {}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        return r.data['disputes'][0]


class AffectedLinesTests(DisputeBase):
    def test_missing_and_damaged_rows_in_bstock_columns(self):
        self._check_in(0, 7)              # 3 not received
        self._check_in(1, 4)
        self._check_in(1, 1, 'salvage')   # 1 damaged
        lines = affected_lines(self.order)
        self.assertEqual(
            [(line['row_number'], line['issue'], line['units']) for line in lines],
            [(1, 'Not received', 3), (2, 'Damaged', 1)],
        )
        first = lines[0]
        self.assertEqual(first['unit_retail'], '50.00')
        self.assertEqual(first['total_retail'], '150.00')
        self.assertEqual(first['upc'], '012345678905')
        self.assertEqual(first['pallet_id'], 'PAL-1')
        self.assertEqual(first['lpn'], 'LPN1')
        self.assertEqual(first['condition'], 'Good')


class StaffApiTests(DisputeBase):
    def test_start_file_and_freeze(self):
        self._check_in(0, 7)
        self._check_in(1, 5)
        d = self._start()
        self.assertEqual(d['status'], 'draft')
        self.assertIn('/dispute/', d['url'])
        self.assertEqual(d['totals']['claimed_retail'], '150.00')
        self.assertIn('3 units ($150.00 manifest retail, 15.0% of the load)', d['suggested_note'])

        url = f'/api/inventory/orders/{self.order.id}/bstock-disputes/{d["id"]}/'
        r = self.client.patch(url, {'status': 'filed', 'bstock_dispute_id': '6392C-D84-418L'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertTrue(r.data['frozen'])
        self.assertIsNotNone(r.data['filed_at'])

        # More units turn up after filing: the filed claim does not move.
        self._check_in(0, 3)
        r = self.client.get(f'/api/inventory/orders/{self.order.id}/bstock-disputes/')
        self.assertEqual(r.data['disputes'][0]['totals']['claimed_retail'], '150.00')

        r = self.client.patch(url, {'status': 'refunded', 'refund_received': '25.50'}, format='json')
        self.assertEqual(r.data['refund_received'], '25.50')
        self.assertIsNotNone(r.data['closed_at'])

        # Back to draft: lines go live again.
        r = self.client.patch(url, {'status': 'draft'}, format='json')
        self.assertEqual(r.data['totals']['claimed_retail'], '0.00')

    def test_bad_status(self):
        d = self._start()
        r = self.client.patch(
            f'/api/inventory/orders/{self.order.id}/bstock-disputes/{d["id"]}/', {'status': 'nope'}, format='json',
        )
        self.assertEqual(r.status_code, 400)


class PublicPageTests(DisputeBase):
    def setUp(self):
        super().setUp()
        self._check_in(0, 7)
        self._check_in(1, 5)
        self.d = self._start()
        self.token = BStockDispute.objects.get(pk=self.d['id']).token
        self.anon = APIClient()

    def test_page_shows_lines_without_sign_in_and_logs_view(self):
        r = self.anon.get(f'/dispute/{self.token}/')
        self.assertEqual(r.status_code, 200)
        html = r.content.decode()
        self.assertIn('TRGET-OXX-1', html)
        self.assertIn('Lotion 1', html)
        self.assertIn('Not received', html)
        self.assertIn('012345678905', html)
        self.assertEqual(r['X-Robots-Tag'], 'noindex, nofollow')
        ev = BStockDisputeEvent.objects.get()
        self.assertEqual(ev.kind, 'view')
        self.assertFalse(ev.signed_in)

    def test_staff_preview_is_marked(self):
        self.anon.get(f'/dispute/{self.token}/?from=dash')
        self.assertTrue(BStockDisputeEvent.objects.get().signed_in)
        r = self.client.get(f'/api/inventory/orders/{self.order.id}/bstock-disputes/')
        self.assertEqual(r.data['disputes'][0]['activity']['outside_views'], 0)

    def test_unknown_token_404(self):
        self.assertEqual(self.anon.get('/dispute/not-a-real-token-123456/').status_code, 404)

    def test_downloads(self):
        r = self.anon.get(f'/dispute/{self.token}/items.csv')
        self.assertEqual(r.status_code, 200)
        self.assertIn('Not received', r.content.decode())

        r = self.anon.get(f'/dispute/{self.token}/dispute.xlsx')
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Items affected', 'Summary', 'Full manifest'])
        self.assertEqual(wb['Items affected']['D2'].value, 'Not received')

        r = self.anon.get(f'/dispute/{self.token}/dispute.zip')
        self.assertEqual(r.status_code, 200)
        zf = zipfile.ZipFile(io.BytesIO(b''.join(r.streaming_content)))
        self.assertEqual(
            sorted(zf.namelist()), ['TRGET-OXX-1-dispute.xlsx', 'TRGET-OXX-1-items-affected.csv'],
        )
        kinds = sorted(BStockDisputeEvent.objects.values_list('kind', flat=True))
        self.assertEqual(kinds, ['csv', 'xlsx', 'zip'])
