"""B-Stock disputes (E2b): the affected lines in B-Stock's own layout, the evidence, and what B-Stock did with it.

B-Stock's dispute form asks the buyer to mark every issue on its on-screen manifest (item, condition, units, MSRP per
unit, total MSRP, UPC) and to attach 1-4 photos per pallet, a truck photo and the BOL (owner's walkthrough,
``.ai/reference/bstock-disputes/``). Dash already holds all of it; this module presents it:

- ``dispute_payload``: order facts, totals and the affected lines (missing = "Not received", broken = "Damaged"), live
  from truck finalize while the dispute is a draft, frozen in ``BStockDispute.snapshot`` once filed.
- The public page ``/dispute/<token>/`` (``views_public_dispute.py``) shows it with the receiving photos, and offers
  Excel, CSV and a photos zip. Every view and download is logged (``BStockDisputeEvent``) so we learn what B-Stock
  actually opens.
"""

from __future__ import annotations

import csv
import hashlib
import io
import secrets
import zipfile
from decimal import Decimal
from typing import Any, Iterator

from django.core.files.storage import default_storage
from django.db import transaction
from django.utils import timezone

from apps.inventory.models import (
    BStockDispute,
    BStockDisputeEvent,
    ManifestRow,
    PurchaseOrder,
    ReceivingAttachment,
    TruckPallet,
)
from apps.inventory.services.truck_finalize import truck_lines, truck_report
from apps.inventory.services.truck_pallets import short_label

ISSUE_MISSING = 'Not received'
ISSUE_BROKEN = 'Damaged'
_CENT = Decimal('0.01')
_SIDE_ORDER = {'front': 0, 'right': 1, 'back': 2, 'left': 3}


def _money(v: Decimal | None) -> str | None:
    return None if v is None else str(Decimal(v).quantize(_CENT))


def _whole(v: Decimal) -> float:
    q = float(v)
    return round(q, 1) if abs(q - round(q)) > 0.05 else float(round(q))


def _condition_label(code: str) -> str:
    return dict(ManifestRow.CONDITION_CHOICES).get(code or '', code or '')


def affected_lines(order: PurchaseOrder) -> list[dict[str, Any]]:
    """One row per (manifest line, issue): units missing are "Not received", units broken are "Damaged"."""
    lines = [line for line in truck_lines(order) if line['units']['missing'] > 0 or line['units']['broken'] > 0]
    meta = {
        mr['id']: mr
        for mr in ManifestRow.objects.filter(pk__in=[line['manifest_row_id'] for line in lines]).values(
            'id', 'condition', 'identifiers', 'tracking', 'brand', 'model',
        )
    }
    out = []
    for line in lines:
        m = meta.get(line['manifest_row_id'], {})
        ids = m.get('identifiers') or {}
        tracking = m.get('tracking') or {}
        for issue, key in ((ISSUE_MISSING, 'missing'), (ISSUE_BROKEN, 'broken')):
            units = line['units'][key]
            if units <= 0:
                continue
            ur = line['unit_retail']
            out.append({
                'row_number': line['row_number'],
                'title': line['title'],
                'condition': _condition_label(m.get('condition') or ''),
                'issue': issue,
                'units': _whole(units),
                'unit_retail': _money(ur),
                'total_retail': _money(units * ur) if ur is not None else None,
                'upc': str(ids.get('upc') or ids.get('UPC') or ''),
                'pallet_id': str(tracking.get('pallet_id') or ''),
                'lpn': str(tracking.get('lpn') or ''),
            })
    return out


def _payload_live(order: PurchaseOrder) -> dict[str, Any]:
    rep = truck_report(order)
    lines = affected_lines(order)
    claimed_units = sum(line['units'] for line in lines)
    claimed_retail = sum((Decimal(line['total_retail']) for line in lines if line['total_retail']), Decimal('0'))
    adjustments = [
        {
            'row_number': a['row_number'], 'title': a['title'], 'units': a['units'], 'timing': a['timing'],
            'note': a['note'],
        }
        for a in rep['adjustments']
    ]
    return {
        'order': {
            'order_number': rep['order']['order_number'],
            'vendor': rep['order']['vendor'],
            'description': rep['order']['description'],
            'delivered_date': rep['order']['delivered_date'],
        },
        'totals': {
            'manifest_units': rep['totals']['units']['manifest'],
            'manifest_retail': rep['totals']['retail']['manifest'],
            'unfulfilled_units': rep['totals']['units']['unfulfilled'],
            'received_units': rep['totals']['units']['received'],
            'missing_units': rep['totals']['units']['missing'],
            'missing_retail': rep['totals']['retail']['missing'],
            'broken_units': rep['totals']['units']['broken'],
            'broken_retail': rep['totals']['retail']['broken'],
            'claimed_units': _whole(Decimal(str(claimed_units))),
            'claimed_retail': _money(claimed_retail),
            'claimed_pct': (
                round(float(claimed_retail) * 100 / float(rep['totals']['retail']['manifest']), 1)
                if rep['totals']['retail']['manifest'] and float(rep['totals']['retail']['manifest'])
                else None
            ),
        },
        'refund_expected': rep['refund']['amount'],
        'adjustments': adjustments,
        'lines': lines,
        'built_at': timezone.now().isoformat(),
    }


def dispute_payload(dispute: BStockDispute) -> dict[str, Any]:
    """Frozen once filed; live while a draft."""
    if dispute.status != BStockDispute.STATUS_DRAFT and dispute.snapshot.get('lines') is not None:
        return dispute.snapshot
    return _payload_live(dispute.purchase_order)


def default_note(payload: dict[str, Any], link: str) -> str:
    t = payload['totals']
    parts = [
        f"{t['claimed_units']:g} units (${t['claimed_retail']} manifest retail"
        + (f", {t['claimed_pct']}% of the load" if t['claimed_pct'] is not None else '')
        + ') were not received or arrived damaged.',
        'Every affected line, the pallet photos, truck photo and BOL are here, with an Excel download: ' + link,
    ]
    return ' '.join(parts)


def create_dispute(user, order: PurchaseOrder) -> BStockDispute:
    with transaction.atomic():
        dispute = BStockDispute.objects.create(
            purchase_order=order,
            token=secrets.token_urlsafe(24),
            created_by=user,
        )
        payload = _payload_live(order)
        dispute.refund_expected = Decimal(payload['refund_expected']) if payload['refund_expected'] else None
        dispute.save(update_fields=['refund_expected', 'updated_at'])
    return dispute


def update_dispute(dispute: BStockDispute, data: dict) -> BStockDispute:
    fields = ['updated_at']
    if 'bstock_dispute_id' in data:
        dispute.bstock_dispute_id = str(data.get('bstock_dispute_id') or '').strip()[:60]
        fields.append('bstock_dispute_id')
    if 'note' in data:
        dispute.note = str(data.get('note') or '')
        fields.append('note')
    if 'refund_received' in data:
        raw = data.get('refund_received')
        if raw in (None, ''):
            dispute.refund_received = None
        else:
            try:
                dispute.refund_received = Decimal(str(raw)).quantize(_CENT)
            except ArithmeticError:
                raise ValueError('Refund received must be a number.') from None
        fields.append('refund_received')
    if 'status' in data:
        new = str(data.get('status') or '')
        if new not in dict(BStockDispute.STATUS_CHOICES):
            raise ValueError('Unknown status.')
        old = dispute.status
        now = timezone.now()
        if new == BStockDispute.STATUS_DRAFT:
            # Back to draft: the lines go live again (a new manifest from B-Stock, a fix on our side).
            dispute.snapshot = {}
            dispute.filed_at = None
            dispute.closed_at = None
        else:
            if old == BStockDispute.STATUS_DRAFT:
                payload = _payload_live(dispute.purchase_order)
                dispute.snapshot = payload
                dispute.refund_expected = Decimal(payload['refund_expected']) if payload['refund_expected'] else None
                dispute.filed_at = now
            dispute.closed_at = now if new in BStockDispute.CLOSED_STATUSES else None
        dispute.status = new
        fields += ['status', 'snapshot', 'filed_at', 'closed_at', 'refund_expected']
    dispute.save(update_fields=list(dict.fromkeys(fields)))
    return dispute


# --- evidence ---------------------------------------------------------------------------------------------------


def dispute_photos(order: PurchaseOrder) -> list[dict[str, Any]]:
    """Receiving photos for the page: BOL, truck, then each pallet's sides with its label number."""
    labels = {
        p.number: short_label(p.code)
        for p in TruckPallet.objects.filter(purchase_order=order).only('number', 'code')
    }
    atts = list(
        ReceivingAttachment.objects.filter(receiving__purchase_order=order)
        .select_related('s3_file', 'thumbnail_file')
    )
    kind_order = {'bol': 0, 'truck': 1, 'pallet_side': 2}
    atts.sort(key=lambda a: (kind_order.get(a.kind, 9), a.pallet_number or 0, _SIDE_ORDER.get(a.side, 9), a.pk))
    out = []
    for a in atts:
        if a.kind == 'bol':
            caption = 'Bill of lading (proof of delivery)'
        elif a.kind == 'truck':
            caption = 'Truck'
        else:
            label = labels.get(a.pallet_number or 0) or ''
            caption = f'Pallet {a.pallet_number}' + (f' ({label})' if label else '') + f' - {a.side or "side"}'
        out.append({
            'id': a.pk,
            'kind': a.kind,
            'pallet_number': a.pallet_number,
            'side': a.side,
            'caption': caption,
            'key': a.s3_file.key if a.s3_file_id else '',
            'url': a.s3_file.url if a.s3_file_id else None,
            'thumb_url': (a.thumbnail_file.url if a.thumbnail_file_id else None) or (a.s3_file.url if a.s3_file_id else None),
        })
    return out


def log_event(dispute: BStockDispute, request, kind: str, detail: str = '') -> None:
    ip = (request.META.get('HTTP_X_FORWARDED_FOR', '') or request.META.get('REMOTE_ADDR', '')).split(',')[0].strip()
    ua = request.META.get('HTTP_USER_AGENT', '') or ''
    visitor = hashlib.sha256(f'{ip}|{ua}'.encode()).hexdigest()[:16]
    BStockDisputeEvent.objects.create(
        dispute=dispute,
        kind=kind,
        detail=detail[:120],
        signed_in=request.GET.get('from') == 'dash',
        visitor=visitor,
        user_agent=ua[:200],
    )


def events_summary(dispute: BStockDispute) -> dict[str, Any]:
    events = list(dispute.events.all()[:200])
    outside = [e for e in events if not e.signed_in]
    by_kind: dict[str, int] = {}
    for e in outside:
        by_kind[e.kind] = by_kind.get(e.kind, 0) + 1
    return {
        'outside_views': by_kind.get(BStockDisputeEvent.KIND_VIEW, 0),
        'outside_by_kind': by_kind,
        'outside_visitors': len({e.visitor for e in outside}),
        'last_outside_at': outside[0].at.isoformat() if outside else None,
        'recent': [
            {
                'kind': e.kind,
                'label': dict(BStockDisputeEvent.KIND_CHOICES).get(e.kind, e.kind),
                'detail': e.detail,
                'staff': e.signed_in,
                'at': e.at.isoformat(),
            }
            for e in events[:25]
        ],
    }


# --- downloads --------------------------------------------------------------------------------------------------

_ITEM_HEADERS = ['Line', 'Item', 'Condition', 'Issue', 'Units', 'MSRP per unit', 'Total MSRP', 'UPC', 'Pallet ID', 'LPN']


def _item_row(line: dict[str, Any]) -> list[Any]:
    return [
        line['row_number'], line['title'], line['condition'], line['issue'], line['units'],
        float(line['unit_retail']) if line['unit_retail'] else None,
        float(line['total_retail']) if line['total_retail'] else None,
        line['upc'], line['pallet_id'], line['lpn'],
    ]


def items_csv(payload: dict[str, Any]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(_ITEM_HEADERS)
    for line in payload['lines']:
        w.writerow(_item_row(line))
    return buf.getvalue()


def workbook_bytes(dispute: BStockDispute, payload: dict[str, Any], link: str) -> bytes:
    """Excel: Items affected (B-Stock's columns), Summary, and the whole manifest with what was received."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = 'Items affected'
    ws.append(_ITEM_HEADERS)
    for line in payload['lines']:
        ws.append(_item_row(line))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip('ABCDEFGHIJ', (6, 60, 14, 14, 8, 14, 14, 16, 34, 22)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = '$#,##0.00'

    t = payload['totals']
    s = wb.create_sheet('Summary')
    for row in (
        ('Order', payload['order']['order_number']),
        ('Seller', payload['order']['vendor']),
        ('Load', payload['order']['description']),
        ('Delivered', payload['order']['delivered_date'] or ''),
        ('B-Stock dispute ID', dispute.bstock_dispute_id),
        ('Manifest units', t['manifest_units']),
        ('Manifest retail', float(t['manifest_retail'] or 0)),
        ('Unfulfilled units (adjusted before/after arrival)', t['unfulfilled_units']),
        ('Received units', t['received_units']),
        ('Not received units', t['missing_units']),
        ('Not received retail', float(t['missing_retail'] or 0)),
        ('Damaged units', t['broken_units']),
        ('Damaged retail', float(t['broken_retail'] or 0)),
        ('Claimed retail', float(t['claimed_retail'] or 0)),
        ('Claimed % of manifest retail', t['claimed_pct']),
        ('Photos and BOL', link),
    ):
        s.append(list(row))
    s.column_dimensions['A'].width = 48
    s.column_dimensions['B'].width = 70

    m = wb.create_sheet('Full manifest')
    m.append(['Line', 'Item', 'Pallet ID', 'MSRP per unit', 'Manifest units', 'Unfulfilled', 'Received',
              'Not received', 'Damaged', 'Over'])
    for line in truck_lines(dispute.purchase_order):
        u = line['units']
        m.append([
            line['row_number'], line['title'], line['pallet_code'],
            float(line['unit_retail']) if line['unit_retail'] is not None else None,
            _whole(u['manifest']), _whole(u['unfulfilled']), _whole(u['received']), _whole(u['missing']),
            _whole(u['broken']), _whole(u['over']),
        ])
    for cell in m[1]:
        cell.font = Font(bold=True)
    m.column_dimensions['B'].width = 60
    m.column_dimensions['C'].width = 34

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class _Pipe(io.RawIOBase):
    """Write-only stream that hands out what was written, so a zip can be streamed as it is built."""

    def __init__(self):
        self.chunks: list[bytes] = []

    def writable(self) -> bool:
        return True

    def write(self, b) -> int:
        self.chunks.append(bytes(b))
        return len(b)

    def take(self) -> bytes:
        data = b''.join(self.chunks)
        self.chunks = []
        return data


def photos_zip_stream(dispute: BStockDispute, payload: dict[str, Any], link: str) -> Iterator[bytes]:
    """Streams a zip: the Excel, the items CSV and every receiving photo at full size."""
    pipe = _Pipe()
    with zipfile.ZipFile(pipe, mode='w', compression=zipfile.ZIP_DEFLATED) as zf:
        prefix = payload['order']['order_number'] or f'order-{dispute.purchase_order_id}'
        zf.writestr(f'{prefix}-dispute.xlsx', workbook_bytes(dispute, payload, link))
        zf.writestr(f'{prefix}-items-affected.csv', items_csv(payload))
        yield pipe.take()
        for i, photo in enumerate(dispute_photos(dispute.purchase_order), start=1):
            if not photo['key']:
                continue
            try:
                with default_storage.open(photo['key'], 'rb') as fh:
                    data = fh.read()
            except Exception:  # noqa: BLE001 - a missing file must not stop the rest
                continue
            name = photo['caption'].replace('/', '-').replace(':', '')
            zf.writestr(f'photos/{i:03d} {name}.jpg', data, compress_type=zipfile.ZIP_STORED)
            yield pipe.take()
    yield pipe.take()


def dispute_out(dispute: BStockDispute, link: str) -> dict[str, Any]:
    """Staff view of one dispute: its fields, the link to send, the claim, and who opened what."""
    payload = dispute_payload(dispute)
    return {
        'id': dispute.pk,
        'status': dispute.status,
        'status_label': dict(BStockDispute.STATUS_CHOICES).get(dispute.status, dispute.status),
        'bstock_dispute_id': dispute.bstock_dispute_id,
        'note': dispute.note,
        'url': link,
        'suggested_note': default_note(payload, link),
        'refund_expected': _money(dispute.refund_expected),
        'refund_received': _money(dispute.refund_received),
        'filed_at': dispute.filed_at.isoformat() if dispute.filed_at else None,
        'closed_at': dispute.closed_at.isoformat() if dispute.closed_at else None,
        'created_at': dispute.created_at.isoformat(),
        'frozen': dispute.status != BStockDispute.STATUS_DRAFT,
        'totals': payload['totals'],
        'line_count': len(payload['lines']),
        'activity': events_summary(dispute),
    }
